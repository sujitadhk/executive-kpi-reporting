"""Build an executive KPI workbook (with native Excel charts) from the consolidated warehouse,
plus CSV exports that can be loaded straight into Power BI."""
import argparse
import re
import sqlite3
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent
DEFAULT_DB = ROOT.parent / "crm-excel-data-consolidation-pipeline" / "data" / "processed" / "warehouse.db"
OUT = ROOT / "output"


def load_queries(path: Path) -> dict[str, str]:
    text = path.read_text()
    parts = re.split(r"-- name: (\w+)\n", text)
    return {parts[i]: parts[i + 1].strip() for i in range(1, len(parts), 2)}


def style_sheet(ws):
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3864")
        cell.alignment = Alignment(horizontal="center")
    for i, col in enumerate(ws.columns, 1):
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col) + 3
        ws.column_dimensions[get_column_letter(i)].width = min(max(width, 12), 34)
    ws.freeze_panes = "A2"


def add_chart(ws, kind, title, cat_col, val_col, anchor, n_rows):
    chart = LineChart() if kind == "line" else BarChart()
    chart.title, chart.height, chart.width = title, 8, 16
    chart.add_data(Reference(ws, min_col=val_col, min_row=1, max_row=n_rows + 1), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=cat_col, min_row=2, max_row=n_rows + 1))
    chart.legend = None
    ws.add_chart(chart, anchor)


def main(db: Path):
    if not db.exists():
        raise SystemExit(f"Warehouse not found at {db}. Run the consolidation pipeline project first.")
    queries = load_queries(ROOT / "sql" / "kpi_queries.sql")
    OUT.mkdir(exist_ok=True)
    with sqlite3.connect(db) as con:
        data = {name: pd.read_sql_query(q, con) for name, q in queries.items()}

    xlsx = OUT / "executive_kpi_report.xlsx"
    with pd.ExcelWriter(xlsx, engine="openpyxl") as xw:
        h = data["headline"].iloc[0]
        summary = pd.DataFrame({"KPI": ["Total revenue", "Total orders", "Active customers",
                                        "Average order value", "Average discount %"],
                                "Value": [h.total_revenue, h.total_orders, h.active_customers,
                                          h.avg_order_value, h.avg_discount_pct]})
        summary.to_excel(xw, sheet_name="Summary", index=False)
        for name in ["monthly_revenue", "revenue_by_region", "revenue_by_category",
                     "revenue_by_segment_country", "top_customers"]:
            data[name].to_excel(xw, sheet_name=name[:31], index=False)
            data[name].to_csv(OUT / f"{name}.csv", index=False)   # Power BI-ready
        summary.to_csv(OUT / "summary_kpis.csv", index=False)

    wb = load_workbook(xlsx)
    for ws in wb.worksheets:
        style_sheet(ws)
    n = len(data["monthly_revenue"])
    add_chart(wb["monthly_revenue"], "line", "Monthly Revenue", 1, 2, "E2", n)
    add_chart(wb["revenue_by_region"], "bar", "Revenue by Region", 1, 2, "F2", len(data["revenue_by_region"]))
    add_chart(wb["revenue_by_category"], "bar", "Revenue by Category", 1, 2, "E2", len(data["revenue_by_category"]))
    for r in range(2, 7):
        wb["Summary"].cell(r, 2).number_format = "#,##0.00"
    wb.save(xlsx)

    top_region = data["revenue_by_region"].iloc[0]
    top_cat = data["revenue_by_category"].iloc[0]
    best_month = data["monthly_revenue"].sort_values("revenue", ascending=False).iloc[0]
    (OUT / "executive_summary.md").write_text(
        f"# Executive Summary\n\n- Total revenue: **{h.total_revenue:,.2f}** from {int(h.total_orders):,} orders "
        f"and {int(h.active_customers)} active customers.\n- Average order value: {h.avg_order_value:,.2f}; "
        f"average discount {h.avg_discount_pct}%.\n- Top region: **{top_region.region}** "
        f"({top_region.revenue:,.2f}).\n- Top category: **{top_cat.category}** ({top_cat.revenue:,.2f}).\n"
        f"- Peak month: **{best_month.month}** ({best_month.revenue:,.2f}).\n")
    print("Report written to", OUT)


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--db", type=Path, default=DEFAULT_DB)
    main(ap.parse_args().db)
